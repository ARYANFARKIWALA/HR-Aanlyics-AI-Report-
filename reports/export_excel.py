"""Corporate Excel Report Generator using openpyxl.

Produces multi-tab formatted workbooks:
- Tab 1: Executive KPI Overview & AI Strategic Narrative
- Tab 2: Formatted Analytical Data Table
- Tab 3: Query Governance, Lineage & Effective-Dating Logic
"""

import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .builder import ReportData


class ExcelReportExporter:
    """Generates styled enterprise HR Excel workbooks."""

    NAVY_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    GRAY_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    HEADER_TITLE = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="1E293B")
    REGULAR_FONT = Font(name="Calibri", size=10, color="1E293B")
    THIN_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    @classmethod
    def generate_excel(cls, report: ReportData, output_path: str | None = None) -> bytes:
        wb = openpyxl.Workbook()

        # ==========================================
        # Sheet 1: Executive Summary & KPIs
        # ==========================================
        ws1 = wb.active
        ws1.title = "Executive Summary"
        ws1.views.sheetView[0].showGridLines = True

        ws1["A1"] = f"HR Analytics: {report.title}"
        ws1["A1"].font = cls.HEADER_TITLE
        ws1["A2"] = f"Generated: {report.generated_at} | Category: {report.category} | Requested By: {report.requested_by}"
        ws1["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

        # KPI Cards Table
        ws1["A4"] = "Enterprise Key Performance Indicators"
        ws1["A4"].font = cls.BOLD_FONT

        kpi_headers = ["Metric", "Value", "Benchmark / Context"]
        for col_idx, h in enumerate(kpi_headers, 1):
            cell = ws1.cell(row=5, column=col_idx, value=h)
            cell.font = cls.WHITE_BOLD
            cell.fill = cls.NAVY_FILL
            cell.alignment = Alignment(horizontal="center")

        kpi_rows = [
            ("Active Headcount", f"{report.kpis.get('active_headcount', 0):,}", "Active Snapshot"),
            ("Annual Turnover Rate", f"{report.kpis.get('attrition_rate_pct', 0)}%", "Voluntary + Involuntary"),
            ("Average Base Salary", f"${report.kpis.get('avg_base_salary', 0):,.2f}", "Across all job grades"),
            ("Average Compa-Ratio", f"{report.kpis.get('avg_compa_ratio', 1.0):.2f}", "Market Midpoint = 1.00"),
            ("Female Representation", f"{report.kpis.get('female_representation_pct', 0)}%", "EEO Diversity Metric"),
            ("Avg Performance Rating", f"{report.kpis.get('avg_performance_rating', 0):.2f} / 5.0", "Latest Annual Cycle"),
        ]

        for r_idx, (m, v, c) in enumerate(kpi_rows, 6):
            ws1.cell(row=r_idx, column=1, value=m).font = cls.BOLD_FONT
            ws1.cell(row=r_idx, column=2, value=v).font = cls.REGULAR_FONT
            ws1.cell(row=r_idx, column=3, value=c).font = cls.REGULAR_FONT
            for c_idx in range(1, 4):
                ws1.cell(row=r_idx, column=c_idx).border = cls.THIN_BORDER

        # AI Narrative
        start_row = 14
        ws1.cell(row=start_row, column=1, value="AI Executive Strategic Narrative").font = cls.BOLD_FONT
        
        narrative_lines = report.executive_summary.split("\n")
        current_r = start_row + 1
        for line in narrative_lines:
            clean_l = line.strip().replace("#", "").replace("*", "")
            if clean_l:
                cell = ws1.cell(row=current_r, column=1, value=clean_l)
                cell.font = cls.REGULAR_FONT
                current_r += 1

        cls._auto_fit_columns(ws1)

        # ==========================================
        # Sheet 2: Detailed Data Table
        # ==========================================
        ws2 = wb.create_sheet(title="Report Data")
        ws2.views.sheetView[0].showGridLines = True

        # Header Row
        for col_idx, col_name in enumerate(report.data_columns, 1):
            cell = ws2.cell(row=1, column=col_idx, value=col_name.replace("_", " ").title())
            cell.font = cls.WHITE_BOLD
            cell.fill = cls.NAVY_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = cls.THIN_BORDER

        # Data Rows
        for row_idx, row_dict in enumerate(report.data_rows, 2):
            fill = cls.ZEBRA_FILL if row_idx % 2 == 0 else PatternFill(fill_type=None)
            for col_idx, col_name in enumerate(report.data_columns, 1):
                val = row_dict.get(col_name, "")
                cell = ws2.cell(row=row_idx, column=col_idx, value=val)
                cell.font = cls.REGULAR_FONT
                cell.border = cls.THIN_BORDER
                if fill.fill_type:
                    cell.fill = fill

        cls._auto_fit_columns(ws2)

        # ==========================================
        # Sheet 3: SQL & Governance
        # ==========================================
        ws3 = wb.create_sheet(title="Lineage & Governance")
        ws3.views.sheetView[0].showGridLines = True

        ws3["A1"] = "Data Lineage, Business Rules & SQL Query"
        ws3["A1"].font = cls.HEADER_TITLE

        ws3["A3"] = "Business Rules Applied"
        ws3["A3"].font = cls.BOLD_FONT
        ws3["A4"] = report.business_rules
        ws3["A4"].font = cls.REGULAR_FONT

        ws3["A6"] = "Effective-Dating Logic"
        ws3["A6"].font = cls.BOLD_FONT
        ws3["A7"] = report.effective_dating_notes
        ws3["A7"].font = cls.REGULAR_FONT

        ws3["A9"] = "Executed SQL Query"
        ws3["A9"].font = cls.BOLD_FONT
        ws3["A10"] = report.sql_query
        ws3["A10"].font = Font(name="Consolas", size=9, color="1E293B")

        cls._auto_fit_columns(ws3)

        buffer = io.BytesIO()
        wb.save(output_path or buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def _auto_fit_columns(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) < 60:
                    max_len = max(max_len, len(val_str))
                else:
                    max_len = max(max_len, 35)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
